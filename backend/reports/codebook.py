"""
Codebook (data dictionary) for every study form, built from the live system.

Nothing here is stored: each download reads
  * the form          -> question order, labels, required
  * the model         -> variable names and types
  * the option tables -> the current choices (value = name)
so a renamed label, a new field or a changed option appears in the next download.
"""
import io
import re

from django import forms as django_forms
from django.db import connection
from django.db.migrations.recorder import MigrationRecorder
from django.template.loader import render_to_string
from django.utils import timezone

from nanopore.forms.diagnosis.diagnosisform import DiagnosisForm
from nanopore.forms.enrollment.enrollmentform import EnrollmentForm
from nanopore.forms.laboratory.clinic.cliniclabform import ClinicLaboratoryForm
from nanopore.forms.laboratory.edcs_tblis.edcs_tblis_lab_form import EdcsTblisZonalForm
from nanopore.forms.laboratory.zonal.zonallabform import ZonalLaboratoryForm
from nanopore.forms.regimen.regimenform import RegimenChangesForm
from nanopore.forms.screening.screeningform import ScreeningForm

# key -> (title, form class). Order = order of the study forms.
CODEBOOK_FORMS = {
    "screening": ("Screening", ScreeningForm),
    "enrollment": ("Enrollment", EnrollmentForm),
    "clinic_lab": ("Clinic Laboratory", ClinicLaboratoryForm),
    "zonal_lab": ("Zonal Laboratory", ZonalLaboratoryForm),
    "edcs_tblis": ("EDCS / TBLIS", EdcsTblisZonalForm),
    "diagnosis": ("Diagnosis", DiagnosisForm),
    "regimen": ("Regimen Changes", RegimenChangesForm),
}

COLUMNS = ["Form", "Q#", "Variable name", "Question label", "Type", "Required", "Values", "Notes"]

# Bookkeeping fields that are not study questions.
SYSTEM_FIELDS = {
    "id", "created_at", "updated_at", "created_by", "updated_by",
    "is_reviewed", "is_locked", "locked_by", "locked_at", "reason_for_change",
}
MAX_LISTED_OPTIONS = 300
QUESTION_NUMBER = re.compile(r"^\s*(\d+[a-z]?(?:\.\d+)?(?:\([a-z]\))?)[.)]?\s+(.*)$", re.S)


def _build_form(form_class):
    try:
        return form_class(screening_instance=None)
    except TypeError:  # forms that don't take a screening
        return form_class()


def _split_number(label):
    """'2. Date sputum received' -> ('2', 'Date sputum received')."""
    match = QUESTION_NUMBER.match(label or "")
    if match:
        return match.group(1), match.group(2).strip()
    return "", (label or "").strip()


def _options(queryset):
    model = queryset.model
    field_names = {f.name for f in model._meta.get_fields()}
    if not {"name", "value"} <= field_names:
        return None  # a link to another record (participant, site, user), not a list of choices
    queryset = queryset.order_by("value", "pk")
    options = list(queryset[: MAX_LISTED_OPTIONS + 1])
    text = ", ".join(f"{o.value}={o.name}" for o in options[:MAX_LISTED_OPTIONS])
    if len(options) > MAX_LISTED_OPTIONS:
        text += ", …"
    return text


def _describe(model_field, form_field):
    """(type, values) for a question."""
    if isinstance(form_field, django_forms.ModelMultipleChoiceField):
        values = _options(form_field.queryset)
        return "Multiple choice (tick all that apply)", values or ""
    if isinstance(form_field, django_forms.ModelChoiceField):
        values = _options(form_field.queryset)
        if values is None:
            return f"Link to {form_field.queryset.model._meta.verbose_name}", ""
        return "Single choice", values
    if isinstance(form_field, django_forms.ChoiceField) and form_field.choices:
        values = ", ".join(f"{k}={v}" for k, v in form_field.choices if k not in ("", None))
        return "Single choice", values

    internal = model_field.get_internal_type() if model_field is not None else ""
    if internal == "DateField" or isinstance(form_field, django_forms.DateField):
        return "Date", "yyyy-mm-dd"
    if internal == "DateTimeField":
        return "Date and time", "yyyy-mm-dd hh:mm"
    if internal in ("IntegerField", "PositiveIntegerField", "SmallIntegerField", "PositiveSmallIntegerField"):
        return "Whole number", ""
    if internal == "DecimalField":
        return f"Number ({model_field.decimal_places} decimal place(s))", ""
    if internal == "FloatField":
        return "Number", ""
    if internal == "BooleanField":
        return "Yes/No (tick box)", "True=Yes, False=No"
    if internal == "TextField" or isinstance(getattr(form_field, "widget", None), django_forms.Textarea):
        return "Free text", ""
    if internal == "CharField":
        return f"Text (max {model_field.max_length} characters)", ""
    return internal or form_field.__class__.__name__, ""


def form_codebook(key):
    """Rows (dicts with COLUMNS keys) for one form, in questionnaire order."""
    title, form_class = CODEBOOK_FORMS[key]
    form = _build_form(form_class)
    model = form._meta.model
    model_fields = {f.name: f for f in model._meta.get_fields() if getattr(f, "concrete", False)}

    rows = []
    for name, form_field in form.fields.items():
        if name in SYSTEM_FIELDS:
            continue
        model_field = model_fields.get(name)
        label = str(form_field.label or (model_field.verbose_name if model_field else name))
        number, label = _split_number(label)
        field_type, values = _describe(model_field, form_field)
        rows.append({
            "Form": title, "Q#": number, "Variable name": name, "Question label": label,
            "Type": field_type, "Required": "Yes" if form_field.required else "No",
            "Values": values, "Notes": str(form_field.help_text or (model_field.help_text if model_field else "") or ""),
        })

    # Variables stored by the system but not asked on the form (e.g. PID, eligibility).
    for name, model_field in model_fields.items():
        if name in form.fields or name in SYSTEM_FIELDS or model_field.many_to_many:
            continue
        field_type, values = _describe(model_field, model_field.formfield() or django_forms.CharField())
        rows.append({
            "Form": title, "Q#": "", "Variable name": name,
            "Question label": str(model_field.verbose_name).capitalize(),
            "Type": field_type, "Required": "", "Values": values,
            "Notes": "Set by the system (not asked on the form)",
        })
    return rows


def codebook_version():
    """Cover information: when it was generated and which database version it describes."""
    recorder = MigrationRecorder(connection)
    latest = (
        recorder.migration_qs.filter(app="nanopore").order_by("-applied", "-id").first()
        if recorder.has_table() else None
    )
    return {
        "generated": timezone.localtime().strftime("%Y-%m-%d %H:%M"),
        "database_version": latest.name if latest else "unknown",
    }


PDF_COLUMNS = COLUMNS[1:]  # each PDF section is one form, so no Form column


def build_codebook(keys):
    forms = []
    for key in keys:
        rows = form_codebook(key)
        forms.append({
            "key": key,
            "title": CODEBOOK_FORMS[key][0],
            "rows": rows,
            "cells": [[row[c] for c in PDF_COLUMNS] for row in rows],
        })
    return {"version": codebook_version(), "forms": forms, "columns": COLUMNS, "pdf_columns": PDF_COLUMNS}


# ── output ───────────────────────────────────────────────────────────────────

def to_excel(codebook):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    cover = workbook.active
    cover.title = "About"
    cover.append(["DREAM Fund – Nanopore codebook"])
    cover["A1"].font = Font(bold=True, size=14)
    cover.append(["Generated", codebook["version"]["generated"]])
    cover.append(["Database version", codebook["version"]["database_version"]])
    cover.append([])
    cover.append(["Form", "Variables"])
    for form in codebook["forms"]:
        cover.append([form["title"], len(form["rows"])])
    cover.column_dimensions["A"].width = 24
    cover.column_dimensions["B"].width = 40

    header_fill = PatternFill("solid", fgColor="1E3A5F")
    widths = {"Form": 18, "Q#": 6, "Variable name": 30, "Question label": 55, "Type": 26,
              "Required": 10, "Values": 60, "Notes": 34}
    for form in codebook["forms"]:
        sheet = workbook.create_sheet(form["title"][:31].replace("/", "-"))
        sheet.append(codebook["columns"])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
        for row in form["rows"]:
            sheet.append([row[c] for c in codebook["columns"]])
        for index, column in enumerate(codebook["columns"], start=1):
            sheet.column_dimensions[sheet.cell(1, index).column_letter].width = widths[column]
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.freeze_panes = "A2"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def to_pdf(codebook):
    from weasyprint import HTML
    html = render_to_string("reports/codebook/codebook_pdf.html", {"codebook": codebook})
    return HTML(string=html).write_pdf()
