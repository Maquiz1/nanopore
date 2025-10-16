# reports/views.py
import io
import json
from django.urls import reverse
from django.shortcuts import render
from django.views import View
from django.db.models import Count
from django.http import HttpResponse
from openpyxl import Workbook
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from django.views.generic import TemplateView
from django.db.models import Count, Q, F
from django.views.generic import ListView
from datetime import date
from dateutil.relativedelta import relativedelta  # pip install python-dateutil
from nanopore.models import Screening, Enrollment, ClinicLaboratory, ZonalLaboratory, Diagnosis
from locations.models import Zone,Site
from weasyprint import HTML

import pdfkit
import openpyxl
from openpyxl.utils import get_column_letter

import pandas as pd
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from django.core.paginator import Paginator
import csv
# from reports.utils import annotate_months_and_substudy  # your helper
from dateutil.relativedelta import relativedelta
from datetime import date


from django.views.generic import TemplateView
from django.db.models import Count, Q
from nanopore.models import Enrollment, Diagnosis, ClinicLaboratory
from django.views.generic import TemplateView
from django.db.models import Count, Q
from nanopore.models import Enrollment, Diagnosis, ClinicLaboratory
from django.template.loader import render_to_string


from django.views.generic import ListView
from django.core.paginator import Paginator
from django.urls import reverse

    
def annotate_months_and_substudy(qs, months_filter=None, substudy_filter=None):
    annotated_qs = []
    for obj in qs:
        # --- Months since TB treatment ---
        if obj.tb_treatment_date:
            delta = relativedelta(date.today(), obj.tb_treatment_date)
            obj.months_since_screening = delta.years * 12 + delta.months
        else:
            obj.months_since_screening = None

        # --- Determine Substudy from ClinicLaboratory ---
        lab = getattr(obj.screening, "clinic_laboratory", None)  # Safe access to one-to-one relation
        if lab and lab.xpert_mtb:
            xpert_id = lab.xpert_mtb.id if hasattr(lab.xpert_mtb, 'id') else lab.xpert_mtb
            if xpert_id in [2, 3, 4, 5, 6]:
                obj.substudy = "Substudy 2"
            elif xpert_id in [1, 7, 8, 9]:
                obj.substudy = "Substudy 4"
            else:
                obj.substudy = "Uncategorized"
        else:
            obj.substudy = "Uncategorized"

        # --- Apply filters ---
        if months_filter and (obj.months_since_screening is None or obj.months_since_screening < months_filter):
            continue
        if substudy_filter and obj.substudy != substudy_filter:
            continue

        annotated_qs.append(obj)

    return annotated_qs


def annotate_months_and_substudy_for_export(qs, months_filter=None, substudy_filter=None):
    annotated_qs = []
    for obj in qs:
        # --- Months since TB treatment ---
        if obj.tb_treatment_date:
            delta = relativedelta(date.today(), obj.tb_treatment_date)
            obj.months_since_treatment = delta.years * 12 + delta.months
        else:
            obj.months_since_treatment = None

        # --- Determine Substudy from ClinicLaboratory ---
        lab = getattr(obj.screening, "clinic_laboratory", None)
        if lab and lab.xpert_mtb:
            xpert_id = lab.xpert_mtb.id if hasattr(lab.xpert_mtb, 'id') else lab.xpert_mtb
            if xpert_id in [2, 3, 4, 5, 6]:
                obj.substudy = "Substudy 2"
            elif xpert_id in [1, 7, 8, 9]:
                obj.substudy = "Substudy 4"
            else:
                obj.substudy = "Uncategorized"
        else:
            obj.substudy = "Uncategorized"

        # --- Apply filters ---
        if months_filter and (obj.months_since_treatment is None or obj.months_since_treatment < months_filter):
            continue
        if substudy_filter and obj.substudy != substudy_filter:
            continue

        annotated_qs.append(obj)
    return annotated_qs


class BaseRecordsListView(ListView):
    template_name = "reports/records_list.html"
    context_object_name = "records"
    model = Diagnosis
    paginate_by = 25
    outcome_filter = None
    exclude_filter = None

    def get_queryset(self):
        qs = Diagnosis.objects.select_related("screening__site__district__region__zone")

        # --- Common GET filters ---
        zone_id = self.request.GET.get("zone")
        site_id = self.request.GET.get("site")
        months_filter = self.request.GET.get("months")
        months_filter = int(months_filter) if months_filter else None
        substudy_filter = self.request.GET.get("substudy")
        start_date = self.request.GET.get("start_date")
        end_date = self.request.GET.get("end_date")

        if zone_id:
            qs = qs.filter(screening__site__district__region__zone_id=zone_id)
        if site_id:
            qs = qs.filter(screening__site_id=site_id)
        if start_date and end_date:
            qs = qs.filter(tb_treatment_date__range=[start_date, end_date])

        # --- Outcome filters ---
        if self.outcome_filter:
            qs = qs.filter(tb_outcome2__in=self.outcome_filter).exclude(tb_outcome2__isnull=True)
        elif self.exclude_filter:
            qs = qs.exclude(tb_outcome2__in=self.exclude_filter)

        return annotate_months_and_substudy(qs, months_filter, substudy_filter)


    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        records_list = self.get_queryset()
        paginator = Paginator(records_list, self.paginate_by)
        page_number = self.request.GET.get("page")
        context["records"] = paginator.get_page(page_number)

        # Export URLs
        query_string = self.request.GET.urlencode()
        context["export_url_excel"] = f"{reverse(self.export_url_name, args=['excel'])}?{query_string}"
        context["export_url_pdf"] = f"{reverse(self.export_url_name, args=['pdf'])}?{query_string}"

        # Common filters
        context["title"] = self.title
        context["selected_zone"] = self.request.GET.get("zone", "")
        context["selected_site"] = self.request.GET.get("site", "")
        context["selected_months"] = self.request.GET.get("months", "")
        context["selected_substudy"] = self.request.GET.get("substudy", "")
        context["zones"] = Zone.objects.all()
        context["sites"] = Site.objects.all()

        # --- Build summaries ---
        substudy2_records = [
            r for r in records_list
            if r.substudy == "Substudy 2"
            and r.months_since_screening is not None
            and r.months_since_screening >= 6
            and not r.tb_outcome2
        ]

        # Per Zone
        zone_summary = {}
        for r in substudy2_records:
            zone_name = r.screening.site.district.region.zone.name
            zone_summary[zone_name] = zone_summary.get(zone_name, 0) + 1

        # Per Site
        site_summary = {}
        for r in substudy2_records:
            site_name = r.screening.site.name
            site_summary[site_name] = site_summary.get(site_name, 0) + 1

        context["zone_summary"] = zone_summary
        context["site_summary"] = site_summary

        return context
    
def substudy_list(request):
    records = Diagnosis.objects.all()  # Or your filtered queryset
    paginator = Paginator(records, 25)  # 25 records per page
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)
    return render(request, 'reports/substudy_list.html', {
        'page_obj': page_obj,
        'title': 'Substudy Records',
        'substudy': request.GET.get('type')
    })
    
# -----------------------
# Dashboard Report (charts)
# -----------------------
def dashboard_report(request):
    visits_per_site = {'Site A': 12, 'Site B': 19, 'Site C': 3, 'Site D': 5}
    mentorship_progress = [3, 7, 4, 6, 8, 10]
    competence_status = {'Completed': 10, 'Pending': 5, 'In Progress': 3}

    context = {
        'visits_labels': json.dumps(list(visits_per_site.keys())),
        'visits_data': json.dumps(list(visits_per_site.values())),
        'progress_labels': json.dumps(['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun']),
        'progress_data': json.dumps(mentorship_progress),
        'status_labels': json.dumps(list(competence_status.keys())),
        'status_data': json.dumps(list(competence_status.values())),
    }
    return render(request, 'reports/index.html', context)


# -----------------------
# Dashboard View
# -----------------------
class ReportDashboardView(View):
    template_name = "reports/dashboard.html"

    def get(self, request, *args, **kwargs):
        return render(request, self.template_name)


# -----------------------
# Summary Report
# -----------------------
from django.views import View
from django.shortcuts import render
from django.db.models import Count

from nanopore.models import Screening, Enrollment

class SummaryReportView(View):
    template_name = "reports/summary.html"

    def get(self, request, *args, **kwargs):
        # Count screenings and enrollments per site/zone
        data = (
            Screening.objects
            .values(
                "site__district__region__zone__name",
                "site__name"
            )
            .annotate(
                total_screenings=Count("id"),
                total_enrollments=Count("enrollment")  # assumes Enrollment has FK 'screening'
            )
            .order_by(
                "site__district__region__zone__name",
                "site__name"
            )
        )

        # Build zone totals + grand totals
        zone_totals = {}
        grand_totals = {"screenings": 0, "enrollments": 0}

        for row in data:
            zone_name = row["site__district__region__zone__name"]

            if zone_name not in zone_totals:
                zone_totals[zone_name] = {"screenings": 0, "enrollments": 0}

            zone_totals[zone_name]["screenings"] += row["total_screenings"]
            zone_totals[zone_name]["enrollments"] += row["total_enrollments"]

            grand_totals["screenings"] += row["total_screenings"]
            grand_totals["enrollments"] += row["total_enrollments"]

        return render(request, self.template_name, {
            "data": data,
            "zone_totals": zone_totals,
            "grand_totals": grand_totals,
        })



# -----------------------
# Forms Report
# -----------------------
class FormsReportView(View):
    template_name = "reports/forms.html"

    def get(self, request, *args, **kwargs):
        zones = Zone.objects.all()
        rows = []
        zone_totals = {}
        grand_totals = {
            "screenings": 0, "enrollments": 0,
            "clinic_lab": 0, "zonal_lab": 0, "diagnosis": 0,
        }

        for zone in zones:
            zone_totals[zone.name] = {
                "screenings": 0, "enrollments": 0,
                "clinic_lab": 0, "zonal_lab": 0, "diagnosis": 0,
            }

            # Get all sites under this zone
            sites = Screening.objects.filter(
                site__district__region__zone=zone
            ).values_list('site', flat=True).distinct()

            from locations.models import Site
            sites = Site.objects.filter(id__in=sites)

            for site in sites:
                row = {
                    "zone": zone.name,
                    "site": site.name,
                    "screenings": Screening.objects.filter(site=site).count(),
                    "enrollments": Enrollment.objects.filter(screening__site=site).count(),
                    "clinic_lab": ClinicLaboratory.objects.filter(screening__site=site).count(),
                    "zonal_lab": ZonalLaboratory.objects.filter(screening__site=site).count(),
                    "diagnosis": Diagnosis.objects.filter(screening__site=site).count(),
                }
                rows.append(row)

                # Add to zone totals
                for k in ["screenings", "enrollments", "clinic_lab", "zonal_lab", "diagnosis"]:
                    zone_totals[zone.name][k] += row[k]
                    grand_totals[k] += row[k]

        return render(request, self.template_name, {
            "rows": rows,
            "zone_totals": zone_totals,
            "grand_totals": grand_totals,
        })


# -----------------------
# Export Summary
# -----------------------
def export_summary(request, fmt):
    # Count screenings and enrollments per site/zone
    data = (
        Screening.objects
        .values(
            "site__district__region__zone__name",
            "site__name"
        )
        .annotate(
            total_screenings=Count("id"),
            total_enrollments=Count("enrollment")  # assumes Enrollment has FK 'screening'
        )
        .order_by(
            "site__district__region__zone__name",
            "site__name"
        )
    )

    # Prepare table data for Excel / PDF
    table_data = [["Zone", "Site", "Total Screenings", "Total Enrollments"]] + [
        [
            row["site__district__region__zone__name"],
            row["site__name"],
            row["total_screenings"],
            row["total_enrollments"],
        ]
        for row in data
    ]

    if fmt == "excel":
        wb = Workbook()
        ws = wb.active
        for r in table_data:
            ws.append(r)
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = 'attachment; filename="summary.xlsx"'
        wb.save(response)
        return response

    elif fmt == "pdf":
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4)
        style = getSampleStyleSheet()
        elements = [Paragraph("Summary Report – Screenings & Enrollments", style["Heading1"])]

        t = Table(table_data)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.lightblue),
            ("GRID", (0, 0), (-1, -1), 1, colors.black),
        ]))
        elements.append(t)
        doc.build(elements)

        pdf = buffer.getvalue()
        buffer.close()
        response = HttpResponse(content_type="application/pdf")
        response["Content-Disposition"] = 'attachment; filename="summary.pdf"'
        response.write(pdf)
        return response



# -----------------------
# Export Forms
# -----------------------
def export_forms(request, fmt):
    table_data = [["Zone", "Site", "Screenings", "Enrollments", "Clinic Lab", "Zonal Lab", "Diagnosis"]]

    zones = Zone.objects.all()
    for zone in zones:
        # Get all sites under this zone
        sites = Screening.objects.filter(
            site__district__region__zone=zone
        ).values_list('site', flat=True).distinct()

        from locations.models import Site
        sites = Site.objects.filter(id__in=sites)

        for site in sites:
            table_data.append([
                zone.name,
                site.name,
                Screening.objects.filter(site=site).count(),
                Enrollment.objects.filter(screening__site=site).count(),
                ClinicLaboratory.objects.filter(screening__site=site).count(),
                ZonalLaboratory.objects.filter(screening__site=site).count(),
                Diagnosis.objects.filter(screening__site=site).count(),
            ])

    if fmt == "excel":
        wb = Workbook()
        ws = wb.active
        for r in table_data:
            ws.append(r)
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = 'attachment; filename="forms.xlsx"'
        wb.save(response)
        return response

    elif fmt == "pdf":
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4)
        style = getSampleStyleSheet()
        elements = [Paragraph("Forms Report", style["Heading1"])]

        t = Table(table_data)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
            ("GRID", (0, 0), (-1, -1), 1, colors.black),
        ]))
        elements.append(t)
        doc.build(elements)

        pdf = buffer.getvalue()
        buffer.close()
        response = HttpResponse(content_type="application/pdf")
        response["Content-Disposition"] = 'attachment; filename="forms.pdf"'
        response.write(pdf)
        return response


class EnrollmentSummaryView(TemplateView):
    template_name = "reports/enrollment_summary.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)

        # --- Base Query ---
        enrollments = Enrollment.objects.select_related(
            "screening__site__district__region__zone"
        )
        labs = ClinicLaboratory.objects.select_related(
            "screening__site__district__region__zone"
        )

        # --- Filters ---
        zone_id = self.request.GET.get("zone")
        site_id = self.request.GET.get("site")
        start_date = self.request.GET.get("start_date")
        end_date = self.request.GET.get("end_date")

        if zone_id:
            enrollments = enrollments.filter(screening__site__district__region__zone_id=zone_id)
            labs = labs.filter(screening__site__district__region__zone_id=zone_id)
        if site_id:
            enrollments = enrollments.filter(screening__site_id=site_id)
            labs = labs.filter(screening__site_id=site_id)
        if start_date and end_date:
            enrollments = enrollments.filter(screening__screening_date__range=[start_date, end_date])
            labs = labs.filter(screening__screening_date__range=[start_date, end_date])

        # --- Aggregations ---
        total_enrolled = enrollments.count()
        substudy2_count = labs.filter(xpert_mtb__in=[2,3,4,5,6]).count()
        substudy4_count = labs.filter(xpert_mtb__in=[1,7,8,9]).count()
        uncategorized_count = labs.exclude(xpert_mtb__in=[1,2,3,4,5,6,7,8,9]).count()

        # --- By Zone ---
        zone_summary = (
            enrollments.values("screening__site__district__region__zone__name")
            .annotate(
                enrolled=Count("id", distinct=True),
                substudy2=Count(
                    "screening__clinic_laboratory",
                    filter=Q(screening__clinic_laboratory__xpert_mtb__in=[2,3,4,5,6]),
                    distinct=True,
                ),
                substudy4=Count(
                    "screening__clinic_laboratory",
                    filter=Q(screening__clinic_laboratory__xpert_mtb__in=[1,7,8,9]),
                    distinct=True,
                ),
                uncategorized=Count(
                    "screening__clinic_laboratory",
                    filter=~Q(screening__clinic_laboratory__xpert_mtb__in=[1,2,3,4,5,6,7,8,9]),
                    distinct=True,
                )
            )
            .order_by("screening__site__district__region__zone__name")
        )

        # --- By Site ---
        site_summary = (
            enrollments.values(
                "screening__site__district__region__zone__name",
                "screening__site__name"
            )
            .annotate(
                enrolled=Count("id", distinct=True),
                substudy2=Count(
                    "screening__clinic_laboratory",
                    filter=Q(screening__clinic_laboratory__xpert_mtb__in=[2,3,4,5,6]),
                    distinct=True,
                ),
                substudy4=Count(
                    "screening__clinic_laboratory",
                    filter=Q(screening__clinic_laboratory__xpert_mtb__in=[1,7,8,9]),
                    distinct=True,
                ),
                uncategorized=Count(
                    "screening__clinic_laboratory",
                    filter=~Q(screening__clinic_laboratory__xpert_mtb__in=[1,2,3,4,5,6,7,8,9]),
                    distinct=True,
                )
            )
            .order_by("screening__site__district__region__zone__name", "screening__site__name")
        )

        # --- Context ---
        context.update({
            "total_enrolled": total_enrolled,
            "substudy2_count": substudy2_count,
            "substudy4_count": substudy4_count,
            "uncategorized_count": uncategorized_count,
            "zone_summary": zone_summary,
            "site_summary": site_summary,
        })

        return context
    
class SubstudyDetailView(TemplateView):
    template_name = "reports/substudy_detail.html"

    def get_context_data(self, substudy, **kwargs):
        context = super().get_context_data(**kwargs)

        # --- Determine filter and title ---
        if substudy == "2":
            xpert_filter = [2, 3, 4, 5, 6]
            context["title"] = "Substudy 2 – Detailed List"
        elif substudy == "4":
            xpert_filter = [1, 7, 8, 9]
            context["title"] = "Substudy 4 – Detailed List"
        elif substudy == "uncategorized":
            xpert_filter = None  # Will filter for values NOT in 1-9
            context["title"] = "Uncategorized – Detailed List"
        else:
            xpert_filter = None
            context["title"] = "Detailed List"

        # --- Base query ---
        labs = ClinicLaboratory.objects.select_related(
            "screening__patient", "screening__site__district__region__zone"
        )

        if xpert_filter is not None:
            labs = labs.filter(xpert_mtb__in=xpert_filter)
        elif substudy == "uncategorized":
            labs = labs.exclude(xpert_mtb__in=[1,2,3,4,5,6,7,8,9])

        # --- Apply filters from GET parameters ---
        zone_name = self.request.GET.get("zone")
        site_id = self.request.GET.get("site")

        if zone_name:
            labs = labs.filter(screening__site__district__region__zone__name=zone_name)
        if site_id:
            labs = labs.filter(screening__site_id=site_id)

        context["records"] = labs
        context["substudy"] = substudy  # Pass to template for export buttons
        return context


class SubstudySummaryView(TemplateView):
    template_name = "reports/substudy_summary.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)

        data = (
            ClinicLaboratory.objects
            .select_related("screening__site__district__region__zone")
            .values(
                "screening__site__district__region__zone__name",
                "screening__site__id",
                "screening__site__name"
            )
            .annotate(
                substudy2=Count("id", filter=Q(xpert_mtb__in=[2,3,4,5,6])),
                substudy4=Count("id", filter=Q(xpert_mtb__in=[1,7,8,9])),
                uncategorized=Count("id", filter=~Q(xpert_mtb__in=[1,2,3,4,5,6,7,8,9]))
            )
            .order_by("screening__site__district__region__zone__name", "screening__site__name")
        )

        context["summary"] = data
        return context
    
    
class SubstudyDetailView(TemplateView):
    template_name = "reports/substudy_detail.html"

    def get_context_data(self, substudy, **kwargs):
        context = super().get_context_data(**kwargs)

        # --- Which substudy filter ---
        if substudy == "2":
            labs = ClinicLaboratory.objects.filter(xpert_mtb__in=[2, 3, 4, 5, 6])
            context["title"] = "Substudy 2 – Detailed List"

        elif substudy == "4":
            labs = ClinicLaboratory.objects.filter(xpert_mtb__in=[1, 7, 8, 9])
            context["title"] = "Substudy 4 – Detailed List"

        elif substudy == "uncategorized":
            labs = ClinicLaboratory.objects.exclude(xpert_mtb__in=[1,2,3,4,5,6,7,8,9])
            context["title"] = "Uncategorized – Detailed List"

        else:
            labs = ClinicLaboratory.objects.none()
            context["title"] = "Unknown Substudy"

        # --- Optimize query ---
        labs = labs.select_related(
            "screening__site__district__region__zone"
        )

        # --- Optional filters ---
        zone_name = self.request.GET.get("zone")
        site_id = self.request.GET.get("site")

        if zone_name:
            labs = labs.filter(screening__site__district__region__zone__name=zone_name)
        if site_id:
            labs = labs.filter(screening__site_id=site_id)

        context["records"] = labs
        return context
    
    

class CompletedStudySummaryView(TemplateView):
    template_name = "reports/completed_study_summary.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        diagnoses = Diagnosis.objects.select_related(
            "screening__site__district__region__zone"
        )

        # Optional filters
        zone_id = self.request.GET.get("zone")
        site_id = self.request.GET.get("site")
        start_date = self.request.GET.get("start_date")
        end_date = self.request.GET.get("end_date")

        if zone_id:
            diagnoses = diagnoses.filter(screening__site__district__region__zone_id=zone_id)
        if site_id:
            diagnoses = diagnoses.filter(screening__site_id=site_id)
        if start_date and end_date:
            diagnoses = diagnoses.filter(screening__screening_date__range=[start_date, end_date])

        # --- Aggregation by Zone ---
        zone_summary = (
            diagnoses
            .values(
                "screening__site__district__region__zone__id",
                "screening__site__district__region__zone__name"
            )
            .annotate(
                completed=Count("id", filter=Q(tb_outcome2__in=[1, 2])),
                in_progress=Count("id", filter=~Q(tb_outcome2__in=[1, 2]))
            )
            .order_by("screening__site__district__region__zone__name")
        )

        # Add friendly keys for template
        zone_summary = [
            {
                "zone_id": z["screening__site__district__region__zone__id"],
                "zone_name": z["screening__site__district__region__zone__name"],
                "completed": z["completed"],
                "in_progress": z["in_progress"],
            }
            for z in zone_summary
        ]

        # --- Aggregation by Site ---
        site_summary = (
            diagnoses
            .values(
                "screening__site__id",
                "screening__site__name",
                "screening__site__district__region__zone__id",
                "screening__site__district__region__zone__name",
            )
            .annotate(
                completed=Count("id", filter=Q(tb_outcome2__in=[1, 2])),
                in_progress=Count("id", filter=~Q(tb_outcome2__in=[1, 2]))
            )
            .order_by(
                "screening__site__district__region__zone__name",
                "screening__site__name"
            )
        )

        # Add friendly keys for template
        site_summary = [
            {
                "site_id": s["screening__site__id"],
                "site_name": s["screening__site__name"],
                "zone_id": s["screening__site__district__region__zone__id"],
                "zone_name": s["screening__site__district__region__zone__name"],
                "completed": s["completed"],
                "in_progress": s["in_progress"],
            }
            for s in site_summary
        ]

        # --- Totals ---
        total_completed = diagnoses.filter(tb_outcome2__in=[1, 2]).exclude(tb_outcome2__isnull=True).count()
        total_in_progress = diagnoses.exclude(tb_outcome2__in=[1, 2]).count()

        context.update({
            "zone_summary": zone_summary,
            "site_summary": site_summary,
            "total_completed": total_completed,
            "total_in_progress": total_in_progress,
        })
        
        return context
    
    
    
# # --- All Completed Records ---
# class CompletedRecordsView(ListView):
#     template_name = "reports/records_list.html"
#     context_object_name = "records"
#     model = Diagnosis
#     paginate_by = 25

#     def get_queryset(self):
#         qs = Diagnosis.objects.select_related("screening__site__district__region__zone")
#         zone_id = self.request.GET.get("zone")
#         site_id = self.request.GET.get("site")
#         start_date = self.request.GET.get("start_date")
#         end_date = self.request.GET.get("end_date")

#         if zone_id:
#             qs = qs.filter(screening__site__district__region__zone_id=zone_id)
#         if site_id:
#             qs = qs.filter(screening__site_id=site_id)
#         if start_date and end_date:
#             qs = qs.filter(screening__screening_date__range=[start_date, end_date])

#         qs = qs.filter(tb_outcome2__in=[1, 2]).exclude(tb_outcome2__isnull=True)
#         return annotate_months_since_screening(qs)

#     def get_context_data(self, **kwargs):
#         context = super().get_context_data(**kwargs)
#         context["title"] = "All Completed Records"
#         return context

# # --- All In-Progress Records ---
# class InProgressRecordsView(ListView):
#     template_name = "reports/records_list.html"
#     context_object_name = "records"
#     model = Diagnosis
#     paginate_by = 25

#     def get_queryset(self):
#         qs = Diagnosis.objects.select_related("screening__site__district__region__zone")
#         zone_id = self.request.GET.get("zone")
#         site_id = self.request.GET.get("site")
#         start_date = self.request.GET.get("start_date")
#         end_date = self.request.GET.get("end_date")

#         if zone_id:
#             qs = qs.filter(screening__site__district__region__zone_id=zone_id)
#         if site_id:
#             qs = qs.filter(screening__site_id=site_id)
#         if start_date and end_date:
#             qs = qs.filter(screening__screening_date__range=[start_date, end_date])

#         qs = qs.exclude(tb_outcome2__in=[1, 2])
#         return annotate_months_since_screening(qs)

#     def get_context_data(self, **kwargs):
#         context = super().get_context_data(**kwargs)
#         context["title"] = "All In Progress Records"
#         return context

# # --- Records by Zone ---
# class RecordsByZoneView(ListView):
#     template_name = "reports/records_list.html"
#     context_object_name = "records"
#     model = Diagnosis
#     paginate_by = 25

#     def get_queryset(self):
#         zone_id = self.kwargs.get("zone_id")
#         status = self.kwargs.get("status")  # "completed" or "in-progress"
#         qs = Diagnosis.objects.select_related("screening__site__district__region__zone").filter(
#             screening__site__district__region__zone_id=zone_id
#         )

#         if status == "completed":
#             qs = qs.filter(tb_outcome2__in=[1, 2]).exclude(tb_outcome2__isnull=True)
#         elif status == "in-progress":
#             qs = qs.exclude(tb_outcome2__in=[1, 2])

#         return annotate_months_since_screening(qs)

#     def get_context_data(self, **kwargs):
#         context = super().get_context_data(**kwargs)
#         zone_id = self.kwargs.get("zone_id")
#         status = self.kwargs.get("status")
#         context["title"] = f"{status.replace('_', ' ').title()} Records in Zone {zone_id}"
#         return context

# # --- Records by Site ---
# class RecordsBySiteView(ListView):
#     template_name = "reports/records_list.html"
#     context_object_name = "records"
#     model = Diagnosis
#     paginate_by = 25

#     def get_queryset(self):
#         site_id = self.kwargs.get("site_id")
#         status = self.kwargs.get("status")  # "completed" or "in-progress"
#         qs = Diagnosis.objects.select_related("screening__site__district__region__zone").filter(
#             screening__site_id=site_id
#         )

#         if status == "completed":
#             qs = qs.filter(tb_outcome2__in=[1, 2]).exclude(tb_outcome2__isnull=True)
#         elif status == "in-progress":
#             qs = qs.exclude(tb_outcome2__in=[1, 2])

#         return annotate_months_since_screening(qs)

#     def get_context_data(self, **kwargs):
#         context = super().get_context_data(**kwargs)
#         site_id = self.kwargs.get("site_id")
#         status = self.kwargs.get("status")
#         context["title"] = f"{status.replace('_', ' ').title()} Records at Site {site_id}"
#         return context


# # --- Helper function to annotate months ---
# def annotate_months_since_screening(qs, months_filter=None):
#     annotated_qs = []
#     for obj in qs:
#         if obj.screening.screening_date:
#             delta = relativedelta(date.today(), obj.screening.screening_date)
#             obj.months_since_screening = delta.years * 12 + delta.months
#         else:
#             obj.months_since_screening = None

#         if months_filter:
#             if obj.months_since_screening is not None and obj.months_since_screening >= months_filter:
#                 annotated_qs.append(obj)
#         else:
#             annotated_qs.append(obj)
#     return annotated_qs

# from django.views.generic import ListView
# from django.core.paginator import Paginator
# from dateutil.relativedelta import relativedelta
# from datetime import date
# from .models import Diagnosis, ClinicLaboratory
# from demographic.models import Zone, Site

# --- All Completed Records ---
# --- All Completed Records ---
class CompletedRecordsView(BaseRecordsListView):
    title = "All Completed Records"
    outcome_filter = [1, 2]   # TB outcomes considered completed
    export_url_name = "reports:export_records"


# --- All In-Progress Records ---
class InProgressRecordsView(BaseRecordsListView):
    title = "All In Progress Records"
    exclude_filter = [1, 2]   # Exclude completed
    export_url_name = "reports:export_records"


# --- Records by Zone ---
class RecordsByZoneView(BaseRecordsListView):
    title = "Records by Zone"
    export_url_name = "reports:export_records"

    def get_queryset(self):
        qs = super().get_queryset()
        zone_id = self.kwargs.get("zone_id")
        status = self.kwargs.get("status")

        qs = [obj for obj in qs if obj.screening.site.district.region.zone.id == int(zone_id)]

        if status == "completed":
            qs = [obj for obj in qs if obj.tb_outcome2 in [1, 2]]
        elif status == "in-progress":
            qs = [obj for obj in qs if obj.tb_outcome2 not in [1, 2]]

        return qs


# --- Records by Site ---
class RecordsBySiteView(BaseRecordsListView):
    title = "Records by Site"
    export_url_name = "reports:export_records"

    def get_queryset(self):
        qs = super().get_queryset()
        site_id = self.kwargs.get("site_id")
        status = self.kwargs.get("status")

        qs = [obj for obj in qs if obj.screening.site.id == int(site_id)]

        if status == "completed":
            qs = [obj for obj in qs if obj.tb_outcome2 in [1, 2]]
        elif status == "in-progress":
            qs = [obj for obj in qs if obj.tb_outcome2 not in [1, 2]]

        return qs


# --- Export Records View ---
class ExportRecordsView(View):
    def get(self, request, export_format, *args, **kwargs):
        qs = Diagnosis.objects.select_related("screening__site__district__region__zone")

        # --- Apply GET filters ---
        zone_id = request.GET.get("zone")
        site_id = request.GET.get("site")
        months_filter = request.GET.get("months")
        months_filter = int(months_filter) if months_filter else None
        substudy_filter = request.GET.get("substudy")
        start_date = request.GET.get("start_date")
        end_date = request.GET.get("end_date")
        status = request.GET.get("status")  # completed / in-progress

        if zone_id:
            qs = qs.filter(screening__site__district__region__zone_id=zone_id)
        if site_id:
            qs = qs.filter(screening__site_id=site_id)
        if start_date and end_date:
            qs = qs.filter(tb_treatment_date__range=[start_date, end_date])

        # --- Status filters ---
        if status == "completed":
            qs = qs.filter(tb_outcome2__in=[1, 2]).exclude(tb_outcome2__isnull=True)
        elif status == "in-progress":
            qs = qs.exclude(tb_outcome2__in=[1, 2])

        # --- Annotate months and substudy ---
        qs = annotate_months_and_substudy_for_export(qs, months_filter, substudy_filter)

        # --- Export ---
        if export_format == "excel":
            return self._export_excel(qs)
        elif export_format == "pdf":
            return self._export_pdf(qs)
        return HttpResponse("Invalid export format", status=400)

    def _export_excel(self, qs):
        import openpyxl
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Records"

        headers = ["PID", "Site", "Zone", "TB Outcome", "TB Treatment Date", "Months Since TB Treatment", "Substudy"]
        ws.append(headers)

        for obj in qs:
            ws.append([
                obj.screening.pid,
                obj.screening.site.name,
                obj.screening.site.district.region.zone.name,
                obj.tb_outcome2,
                obj.tb_treatment_date,
                obj.months_since_treatment,
                obj.substudy
            ])

        # Auto-width columns
        for i, col in enumerate(ws.columns, 1):
            max_length = max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[get_column_letter(i)].width = max_length + 2

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = "attachment; filename=records.xlsx"
        wb.save(response)
        return response

    def _export_pdf(self, qs):
        from django.template.loader import render_to_string
        import pdfkit

        # --- Get filters ---
        months_filter = self.request.GET.get("months")
        substudy_filter = self.request.GET.get("substudy")
        status = self.request.GET.get("status")
        zone_id = self.request.GET.get("zone")
        site_id = self.request.GET.get("site")
        start_date = self.request.GET.get("start_date")
        end_date = self.request.GET.get("end_date")

        # --- Build header text ---
        header_parts = []

        # Status label
        if status == "completed":
            header_parts.append("Completed Patients")
        elif status == "in-progress":
            header_parts.append("In-Progress Patients")
        else:
            header_parts.append("All Patients")

        # Months filter
        if months_filter:
            header_parts.append(f"with ≥ {months_filter} Months of Treatment")

        # Substudy filter
        if substudy_filter:
            header_parts.append(f"({substudy_filter})")

        # Zone and Site Names
        zone_name = None
        site_name = None
        if zone_id:
            from locations.models import Zone  # adjust app/model name
            zone = Zone.objects.filter(pk=zone_id).first()
            if zone:
                zone_name = zone.name
                header_parts.append(f"Zone: {zone.name}")

        if site_id:
            from locations.models import Site  # adjust app/model name
            site = Site.objects.filter(pk=site_id).first()
            if site:
                site_name = site.name
                header_parts.append(f"Site: {site.name}")

        # Date range
        if start_date and end_date:
            header_parts.append(f"Period: {start_date} to {end_date}")

        # Final title
        header_text = " - ".join(header_parts) if header_parts else "Patient Records"

        # --- Add numbering ---
        records = []
        for i, obj in enumerate(qs, start=1):
            records.append({
                "no": i,
                "pid": obj.screening.pid,
                "site": obj.screening.site.name,
                "zone": obj.screening.site.district.region.zone.name,
                "tb_outcome": obj.tb_outcome2,
                "treatment_date": obj.tb_treatment_date,
                "months_since_treatment": obj.months_since_treatment,
                "substudy": obj.substudy,
            })

        total_records = len(records)

        # --- Render HTML template ---
        html = render_to_string("reports/records_export_pdf.html", {
            "header_text": header_text,
            "records": records,
            "total_records": total_records,
            "zone_name": zone_name,
            "site_name": site_name,
            "start_date": start_date,
            "end_date": end_date,
        })

        # --- PDF with page numbers ---
        options = {
            "page-size": "A4",
            "encoding": "UTF-8",
            "footer-right": "Page [page] of [toPage]",
            "footer-font-size": "9",
            "margin-bottom": "15mm",
        }

        pdf = pdfkit.from_string(html, False, options=options)

        response = HttpResponse(pdf, content_type="application/pdf")
        response["Content-Disposition"] = "attachment; filename=records.pdf"
        return response


    
        
class ScreeningSummaryView(TemplateView):
    template_name = "reports/screening_summary.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)

        # --- Base Query ---
        screenings = Screening.objects.select_related(
            "site__district__region__zone"
        )

        # --- Filters from request ---
        zone_id = self.request.GET.get("zone")
        site_id = self.request.GET.get("site")
        start_date = self.request.GET.get("start_date")
        end_date = self.request.GET.get("end_date")

        if zone_id:
            screenings = screenings.filter(site__district__region__zone_id=zone_id)
        if site_id:
            screenings = screenings.filter(site_id=site_id)
        if start_date and end_date:
            screenings = screenings.filter(screening_date__range=[start_date, end_date])

        # --- Aggregations ---
        total_screenings = screenings.count()
        total_not_eligible = screenings.filter(eligible=False).count()

        # --- Summary by Zone ---
        zone_summary = (
            screenings.values("site__district__region__zone__name")
            .annotate(
                total_screenings=Count("id", distinct=True),
                not_eligible_count=Count("id", filter=Q(eligible=False), distinct=True),
            )
            .order_by("site__district__region__zone__name")
        )

        # --- Summary by Site ---
        site_summary = (
            screenings.values(
                "site__district__region__zone__name",
                "site__name"
            )
            .annotate(
                total_screenings=Count("id", distinct=True),
                not_eligible_count=Count("id", filter=Q(eligible=False), distinct=True),
            )
            .order_by("site__district__region__zone__name", "site__name")
        )

        # --- Zone Totals (for table footer) ---
        zone_totals = {}
        for zone in zone_summary:
            zone_name = zone["site__district__region__zone__name"]
            zone_totals[zone_name] = {
                "screenings": zone["total_screenings"],
                "not_eligible": zone["not_eligible_count"],
            }

        # --- Grand Totals ---
        grand_totals = {
            "screenings": total_screenings,
            "not_eligible": total_not_eligible,
        }

        context.update({
            "data": site_summary,
            "zone_totals": zone_totals,
            "grand_totals": grand_totals,
        })

        return context


  
class EligibilitySummaryView(TemplateView):
    template_name = "reports/eligibility_summary.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)

        # --- Base Query ---
        screenings = Screening.objects.select_related(
            "site__district__region__zone"
        )

        # --- Filters from request ---
        zone_id = self.request.GET.get("zone")
        site_id = self.request.GET.get("site")
        start_date = self.request.GET.get("start_date")
        end_date = self.request.GET.get("end_date")

        if zone_id:
            screenings = screenings.filter(site__district__region__zone_id=zone_id)
        if site_id:
            screenings = screenings.filter(site_id=site_id)
        if start_date and end_date:
            screenings = screenings.filter(screening_date__range=[start_date, end_date])

        # --- Aggregations ---
        total_screenings = screenings.count()
        total_not_eligible = screenings.filter(eligible=False).count()

        # --- Summary by Zone ---
        zone_summary = (
            screenings.values("site__district__region__zone__name")
            .annotate(
                total_screenings=Count("id", distinct=True),
                not_eligible_count=Count("id", filter=Q(eligible=False), distinct=True),
            )
            .order_by("site__district__region__zone__name")
        )

        # --- Summary by Site ---
        site_summary = (
            screenings.values(
                "site__district__region__zone__name",
                "site__name"
            )
            .annotate(
                total_screenings=Count("id", distinct=True),
                not_eligible_count=Count("id", filter=Q(eligible=False), distinct=True),
            )
            .order_by("site__district__region__zone__name", "site__name")
        )

        # --- Zone Totals (for table footer) ---
        zone_totals = {}
        for zone in zone_summary:
            zone_name = zone["site__district__region__zone__name"]
            zone_totals[zone_name] = {
                "screenings": zone["total_screenings"],
                "not_eligible": zone["not_eligible_count"],
            }

        # --- Grand Totals ---
        grand_totals = {
            "screenings": total_screenings,
            "not_eligible": total_not_eligible,
        }

        context.update({
            "data": site_summary,
            "zone_totals": zone_totals,
            "grand_totals": grand_totals,
        })

        return context
    
    
class SubstudyExportView(View):
    def get(self, request, export_format):
        # Get type from query params: "2", "4", or "uncategorized"
        export_type = request.GET.get("type")  

        # --- Determine filter and title ---
        if export_type == "2":
            xpert_filter = [2, 3, 4, 5, 6]
            title = "Substudy 2 Records"
        elif export_type == "4":
            xpert_filter = [1, 7, 8, 9]
            title = "Substudy 4 Records"
        elif export_type == "uncategorized":
            xpert_filter = None  # will exclude Substudy 2 and 4
            title = "Uncategorized Records"
        else:
            return HttpResponse("Invalid type", status=400)

        # --- Base query ---
        labs = ClinicLaboratory.objects.select_related(
            "screening__site__district__region__zone"
        )

        # --- Filter by type ---
        if xpert_filter is not None:
            labs = labs.filter(xpert_mtb__in=xpert_filter)
        else:
            labs = labs.exclude(xpert_mtb__in=[1,2,3,4,5,6,7,8,9])

        # --- Optional filters from GET ---
        zone_name = request.GET.get("zone")
        site_id = request.GET.get("site")
        if zone_name:
            labs = labs.filter(screening__site__district__region__zone__name=zone_name)
        if site_id:
            labs = labs.filter(screening__site__id=site_id)

        # --- Prepare dataset ---
        data = [
            {
                "Zone": lab.screening.site.district.region.zone.name,
                "Site": lab.screening.site.name,
                "Screening ID": lab.screening.pid,
                "Screening Date": lab.screening.screening_date,
                "Xpert MTB Result": lab.xpert_mtb,
            }
            for lab in labs
        ]

        # --- Export to Excel ---
        if export_format == "excel":
            df = pd.DataFrame(data)
            buffer = io.BytesIO()
            with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
                df.to_excel(writer, index=False, sheet_name="Substudy")
            buffer.seek(0)
            response = HttpResponse(
                buffer,
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            response["Content-Disposition"] = f'attachment; filename="{title}.xlsx"'
            return response

        # --- Export to PDF ---
        elif export_format == "pdf":
            buffer = io.BytesIO()
            p = canvas.Canvas(buffer, pagesize=A4)
            width, height = A4

            p.setFont("Helvetica-Bold", 14)
            p.drawString(50, height - 50, title)

            p.setFont("Helvetica", 10)
            y = height - 80
            for row in data:
                line = f"{row['Zone']} | {row['Site']} | {row['Screening ID']} | {row['Screening Date']} | {row['Xpert MTB Result']}"
                p.drawString(50, y, line)
                y -= 15
                if y < 50:
                    p.showPage()
                    y = height - 50

            p.save()
            buffer.seek(0)
            response = HttpResponse(buffer, content_type="application/pdf")
            response["Content-Disposition"] = f'attachment; filename="{title}.pdf"'
            return response

        return HttpResponse("Invalid format", status=400)