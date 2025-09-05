from django.views import View
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from nanopore.models import Screening

class ScreeningExcelDownloadView(View):
    """
    Download screening records as an Excel file, with formatting.
    Filters based on user's allowed sites or zones.
    """
    def get(self, request, *args, **kwargs):
        # Create workbook and worksheet
        wb = Workbook()
        ws = wb.active
        ws.title = "Screenings"

        # Define header
        headers = [
            'PID', 'Screening Date', 'Sex', 'DOB', 'Age', 'Present Symptoms',
            'Genexpert Confirmation', 'Produce Respiratory Sample', 'Age >=18',
            'Consent', 'Consent Date', 'Not Willing', 'Unable to Understand',
            'Enrolled', 'Reason', 'Other Reason', 'Remarks', 'Eligible', 'Site', 'Zone'
        ]

        # Write header with bold font
        bold_font = Font(bold=True)
        for col_num, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.font = bold_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        user = request.user
        screenings = Screening.objects.all()

        # Filter by user's allowed sites/zones if applicable
        if hasattr(user, 'sites') and user.sites.exists():
            screenings = screenings.filter(site__in=user.sites.all())
        if hasattr(user, 'zones') and user.zones.exists():
            screenings = screenings.filter(site__district__region__zone__in=user.zones.all())
        if user.is_superuser:
            screenings = Screening.objects.all()

        screenings = screenings.order_by('screening_date')

        # Write data rows
        for row_num, s in enumerate(screenings, start=2):
            zone_name = ""
            if s.site and s.site.district and s.site.district.region and s.site.district.region.zone:
                zone_name = s.site.district.region.zone.name

            row = [
                s.pid,
                s.screening_date,
                s.sex.name if s.sex else '',
                s.dob,
                s.age,
                s.present_symptoms.name if s.present_symptoms else '',
                s.genexpert_confirmation.name if s.genexpert_confirmation else '',
                s.produce_resp_sample.name if s.produce_resp_sample else '',
                s.age18years.name if s.age18years else '',
                s.consent.name if s.consent else '',
                s.consent_date,
                s.not_willing.name if s.not_willing else '',
                s.unable_understand.name if s.unable_understand else '',
                s.enrolled.name if s.enrolled else '',
                s.reasons.name if s.reasons else '',
                s.reasons_other or '',
                s.remarks or '',
                s.eligible,
                s.site.name if s.site else '',
                zone_name,
            ]
            for col_num, value in enumerate(row, start=1):
                cell = ws.cell(row=row_num, column=col_num, value=value)
                cell.alignment = Alignment(horizontal="center", vertical="center")

        # Adjust column widths
        for column_cells in ws.columns:
            max_length = max((len(str(cell.value)) if cell.value else 0) for cell in column_cells)
            adjusted_width = max_length + 2
            ws.column_dimensions[column_cells[0].column_letter].width = adjusted_width

        # Prepare response
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="screenings.xlsx"'
        wb.save(response)
        return response
